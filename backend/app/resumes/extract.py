"""文本提取：pypdf / python-docx，纯文本型文件。

- 扫描件/图片型 PDF 提不出文本 → NoTextLayerError（安全失败，不伪装成功）。
- 不做 OCR（阶段边界）。
"""

import io

from docx import Document
from pypdf import PdfReader

from app.resumes.validation import DOCX_MIME, PDF_MIME

# 低于该字符数视为没有可用文本层（扫描件或空文档）
MIN_TEXT_CHARS = 30


class NoTextLayerError(Exception):
    """文件没有可提取的文本层（扫描件 / 图片型 PDF / 空文档）。"""


class ExtractionError(Exception):
    """提取过程失败（损坏、格式异常等）。"""


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        parts = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ExtractionError("pdf extraction failed") from exc
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    except Exception as exc:
        raise ExtractionError("docx extraction failed") from exc
    return "\n".join(parts)


def extract_text(data: bytes, media_type: str) -> str:
    """按媒体类型提取纯文本；文本层不足抛 NoTextLayerError。"""
    if media_type == PDF_MIME:
        text = _extract_pdf(data)
    elif media_type == DOCX_MIME:
        text = _extract_docx(data)
    else:
        raise ExtractionError(f"unsupported media type: {media_type}")

    if len(text.strip()) < MIN_TEXT_CHARS:
        raise NoTextLayerError("no usable text layer")
    return text
