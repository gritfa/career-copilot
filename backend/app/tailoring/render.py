"""标准单栏模板渲染（ADR-001：只做 1 套模板，多模板推迟）。

选型（纯 Python，无系统级依赖）：
- DOCX：python-docx（已是主依赖；可重新打开编辑）；
- PDF：reportlab + 内置 Adobe CID 中文字体 STSong-Light
  （无需字体文件/libpango 等系统依赖；文本层真实可选择/复制）。

约束：
- 渲染只消费已通过确定性校验的 ResumeContent，本层不做任何内容改写；
- 长文本自动折行、跨页续排，不截断正文；
- 本模块不写日志（避免任何简历正文进入日志的可能）。
"""

import io

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from app.tailoring.schemas import ResumeContent

_PDF_FONT = "STSong-Light"
_DOCX_EAST_ASIAN_FONT = "宋体"
_pdf_font_registered = False

_DOC_TITLE = "定制简历"


def _ensure_pdf_font() -> None:
    global _pdf_font_registered
    if not _pdf_font_registered:
        pdfmetrics.registerFont(UnicodeCIDFont(_PDF_FONT))
        _pdf_font_registered = True


def _subtitle(content: ResumeContent) -> str:
    parts = [p for p in (content.target_company, content.target_job_title) if p]
    return f"目标岗位：{' · '.join(parts)}" if parts else ""


def render_resume(content: ResumeContent, fmt: str) -> bytes:
    if fmt == "docx":
        return render_docx(content)
    if fmt == "pdf":
        return render_pdf(content)
    raise ValueError(f"unsupported export format: {fmt}")


# ---------------- DOCX ----------------


def _docx_set_font(run, size: int, bold: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = _DOCX_EAST_ASIAN_FONT
    # 中文字形需要显式设置 eastAsia，否则部分渲染器回退到拉丁字体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _DOCX_EAST_ASIAN_FONT)


def render_docx(content: ResumeContent) -> bytes:
    doc = DocxDocument()
    title = doc.add_paragraph()
    _docx_set_font(title.add_run(_DOC_TITLE), size=18, bold=True)
    subtitle = _subtitle(content)
    if subtitle:
        para = doc.add_paragraph()
        _docx_set_font(para.add_run(subtitle), size=11)

    for section in content.sections:
        heading = doc.add_paragraph()
        _docx_set_font(heading.add_run(section.title), size=14, bold=True)
        if not section.items:
            para = doc.add_paragraph()
            _docx_set_font(para.add_run("（本节暂无已确认内容）"), size=10)
            continue
        for item in section.items:
            para = doc.add_paragraph(style="List Bullet")
            _docx_set_font(para.add_run(item.text), size=11)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------

_PAGE_W, _PAGE_H = A4
_MARGIN_X = 50
_MARGIN_TOP = 60
_MARGIN_BOTTOM = 60
_BODY_SIZE = 11
_LINE_HEIGHT = 18


def _wrap_line(text: str, size: int, max_width: float) -> list[str]:
    """按实际字宽折行（中文无空格分词，逐字符累积）。"""
    lines: list[str] = []
    current = ""
    for ch in text:
        candidate = current + ch
        if pdfmetrics.stringWidth(candidate, _PDF_FONT, size) > max_width and current:
            lines.append(current)
            current = ch
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines or [""]


class _PdfWriter:
    """极简单栏排版器：自动折行 + 跨页续排。"""

    def __init__(self) -> None:
        self.buf = io.BytesIO()
        self.canvas = canvas.Canvas(self.buf, pagesize=A4)
        self.y = _PAGE_H - _MARGIN_TOP

    def _ensure_room(self, needed: float) -> None:
        if self.y - needed < _MARGIN_BOTTOM:
            self.canvas.showPage()
            self.y = _PAGE_H - _MARGIN_TOP

    def text(self, text: str, size: int, indent: float = 0, gap: float = 0) -> None:
        max_width = _PAGE_W - 2 * _MARGIN_X - indent
        for line in _wrap_line(text, size, max_width):
            self._ensure_room(_LINE_HEIGHT)
            self.canvas.setFont(_PDF_FONT, size)
            self.canvas.drawString(_MARGIN_X + indent, self.y - size, line)
            self.y -= _LINE_HEIGHT
        self.y -= gap

    def finish(self) -> bytes:
        self.canvas.showPage()
        self.canvas.save()
        return self.buf.getvalue()


def render_pdf(content: ResumeContent) -> bytes:
    _ensure_pdf_font()
    writer = _PdfWriter()
    writer.text(_DOC_TITLE, size=18, gap=4)
    subtitle = _subtitle(content)
    if subtitle:
        writer.text(subtitle, size=_BODY_SIZE, gap=8)

    for section in content.sections:
        writer.text(section.title, size=14, gap=2)
        if not section.items:
            writer.text("（本节暂无已确认内容）", size=10, indent=12, gap=6)
            continue
        for item in section.items:
            writer.text(f"· {item.text}", size=_BODY_SIZE, indent=12)
        writer.y -= 6
    return writer.finish()
